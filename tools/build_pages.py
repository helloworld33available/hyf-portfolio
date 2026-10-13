# -*- coding: utf-8 -*-
"""把 md / docx 源稿转成作品集详情页。

用法：
    python tools/build_pages.py

产出全部在 works/ 下。每篇的详细 meta 在下面的 WORKS 列表里维护。
"""
import os
import re
import html
import base64
import shutil

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # portfolio/
SRC = os.path.dirname(ROOT)                                          # 工作区根目录（md 源稿）
DESKTOP = r'C:\Users\萤火石之眼\Desktop\嘟嘟哒\游戏策划'
OUT = os.path.join(ROOT, 'works')
IMG = os.path.join(OUT, 'assets', 'img')

import markdown  # noqa: E402
from docx import Document  # noqa: E402
from docx.table import Table  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

WORKS = [
    dict(
        slug='e7-monetization-design',
        title='《第七史诗》商业化系统策划案 · 回声铭牌 EM',
        cat='design',
        cat_label='系统策划案',
        role='系统策划 · 商业化方向',
        tag='tag--design',
        year='2026.10',
        metric_a=('V0.2', '文档版本'),
        metric_b=('5%', '额外产出上限'),
        source=os.path.join(SRC, '第七史诗_商业化系统策划案.md'),
        kind='md',
    ),
    dict(
        slug='hsr-warp-design',
        title='《崩坏：星穹铁道》跃迁系统策划案',
        cat='design',
        cat_label='系统策划案',
        role='系统策划 · 抽卡与保底',
        tag='tag--design',
        year='2026.09',
        metric_a=('v1.2', '文档版本'),
        metric_b=('13', '界面原型图'),
        source=os.path.join(DESKTOP, '策划案', '《崩坏：星穹铁道》跃迁系统策划案1.2.docx'),
        kind='docx',
    ),
    dict(
        slug='luoke-world-design',
        title='《洛克王国：世界》通行证系统策划案',
        cat='design',
        cat_label='系统策划案',
        role='系统策划 · 赛季通行证',
        tag='tag--design',
        year='2026.05',
        metric_a=('V0.4', '文档版本'),
        metric_b=('21', '界面与配图'),
        source=os.path.join(DESKTOP, '策划案', '《洛克王国：世界》.docx'),
        kind='docx',
        skip_head=3,
    ),
    dict(
        slug='e7-economy-analysis',
        title='《第七史诗》经济系统拆解案',
        cat='analysis',
        cat_label='系统拆解案',
        role='拆解分析 · 资源流动',
        tag='tag--analysis',
        year='2026.10',
        metric_a=('31', '表格数'),
        metric_b=('4', '核心结论'),
        source=os.path.join(SRC, '第七史诗_经济系统拆解案.md'),
        kind='md',
    ),
    dict(
        slug='e7-monetization-analysis',
        title='《第七史诗》商业化系统拆解案',
        cat='analysis',
        cat_label='系统拆解案',
        role='拆解分析 · 付费全链路',
        tag='tag--analysis',
        year='2026.10',
        metric_a=('Φ', '基准单位'),
        metric_b=('7', '付费板块'),
        source=os.path.join(SRC, '第七史诗_商业化系统拆解案.md'),
        kind='md',
    ),
    dict(
        slug='e7-payment-analysis',
        title='《第七史诗》付费系统拆解案',
        cat='analysis',
        cat_label='系统拆解案',
        role='拆解分析 · 商城与付费分层',
        tag='tag--analysis',
        year='2026.10',
        metric_a=('15', '表格数'),
        metric_b=('121', '抽硬保底'),
        source=os.path.join(DESKTOP, '拆解案', '《第七史诗》付费系统拆解案.docx'),
        kind='docx',
    ),
    dict(
        slug='hsr-payment-analysis',
        title='《崩坏：星穹铁道》付费系统拆解案',
        cat='analysis',
        cat_label='系统拆解案',
        role='拆解分析 · 付费与市场',
        tag='tag--analysis',
        year='2026.10',
        metric_a=('v3', '文档版本'),
        metric_b=('10', '表格数'),
        source=os.path.join(DESKTOP, '拆解案', '《崩坏：星穹铁道》付费系统拆解案_v3.docx'),
        kind='docx',
    ),
    dict(
        slug='wz-gift-analysis',
        title='《王者荣耀》商城礼包体系拆解案',
        cat='analysis',
        cat_label='系统拆解案',
        role='拆解分析 · 礼包与定价',
        tag='tag--analysis',
        year='2026.10',
        metric_a=('9', '礼包层级'),
        metric_b=('3', '定价模型'),
        source=os.path.join(SRC, '《王者荣耀》商城礼包体系拆解案.md'),
        kind='md',
    ),
    dict(
        slug='ares-combat-analysis',
        title='《阿瑞斯病毒》战斗系统拆解案',
        cat='analysis',
        cat_label='系统拆解案',
        role='拆解分析 · 战斗手感',
        tag='tag--analysis',
        year='2026.10',
        metric_a=('4', '拆解维度'),
        metric_b=('2', '战斗层次'),
        source=os.path.join(SRC, '阿瑞斯病毒_战斗系统拆解案.md'),
        kind='md',
    ),
    dict(
        slug='oasis-xiuxian-analysis',
        title='《和平精英·绿洲启元》吃丹修仙玩法拆解案',
        cat='analysis',
        cat_label='系统拆解案',
        role='拆解分析 · UGC 玩法',
        tag='tag--analysis',
        year='2026.10',
        metric_a=('6', '修订版本'),
        metric_b=('5', '成长阶段'),
        source=os.path.join(DESKTOP, '拆解案', '《和平精英·绿洲启元》吃丹修仙玩法拆解案.docx'),
        kind='docx',
    ),
]

TPL = u'''<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title} · 黄彦飞</title>
<meta name="description" content="{title} — 系统策划作品集 by 黄彦飞">
<link rel="stylesheet" href="../assets/styles.css">
</head>
<body>
<nav class="nav">
  <div class="nav__inner">
    <a class="nav__logo" href="../index.html">HYF<span>系统策划作品集</span></a>
    <button class="nav__toggle" aria-label="菜单">&#9776;</button>
    <ul class="nav__links">
      <li><a href="../index.html#about"><em>01</em>关于</a></li>
      <li><a href="../index.html#skills"><em>02</em>核心能力</a></li>
      <li><a href="../index.html#work"><em>03</em>代表作品</a></li>
      <li><a href="../index.html#play"><em>04</em>游戏经历</a></li>
      <li><a href="../index.html#contact"><em>05</em>联系方式</a></li>
    </ul>
  </div>
</nav>

<main class="page">
  <div class="page__body">
    <a class="page__back" href="../index.html#work">&#8592; 返回作品列表</a>

    <header class="doc-head">
      <span class="doc-head__tag {tag}">{cat_label}</span>
      <h1>{title}</h1>
      <p class="doc-head__desc">{desc}</p>
      <div class="doc-meta">
        <span>撰写 / {year}</span>
        <span>{role}</span>
        <span>黄彦飞</span>
      </div>
    </header>

    <nav class="doc-nav" id="toc-box" style="display:none">
      <div class="doc-nav__title">目录</div>
      <ol id="toc-list"></ol>
    </nav>

    <article class="doc">
{body}
    </article>

    <div class="doc-foot">
      <a class="page__back" href="../index.html#work">&#8592; 返回作品列表</a>
      <a class="page__back" href="../index.html#contact">联系我 &#8594;</a>
    </div>
  </div>
</main>

<footer class="footer">黄彦飞 · 游戏系统策划作品集 &nbsp;|&nbsp; 上海</footer>
<script src="../assets/main.js"></script>
</body>
</html>
'''

DESCS = {
    'e7-monetization-design': u'针对《第七史诗》长尾付费塌陷，提出「回声铭牌 EM」记账型反馈模块：以重复获取角色的残余价值为锚点重新计价，用资源类与典藏类双目录把回收出口与版本通胀预算隔开，并对肝帝、氪佬、又氪又肝三类人群分别做发行预算校验。含完整配置表结构与核心逻辑伪代码。',
    'hsr-warp-design': u'完整设计一套抽卡（跃迁）系统：卡池结构与 UI 分区、概率与保底递增量表、跨池继承规则、UP 切换与追歼值落的边界处理，配套 13 张界面原型图与状态流转说明。',
    'luoke-world-design': u'围绕《洛克王国：世界》「拾光通行证」做复刻式系统设计：双入口跳转逻辑与首次/非首次进入的分支状态、等级构成与经验投放、任务链与赛季活动的对应关系、免费线与小/大礼券双付费线的奖励点布局，配套界面布局图与配表说明。',
    'e7-economy-analysis': u'从 faucet / sink 视角还原《第七史诗》的资源流动结构，解释这套运营多年的产品为什么能把体力的产出增长率压到接近零，以及金币作为共享稀缺资源如何在低氪玩家的每日决策里制造取舍。',
    'e7-monetization-analysis': u'把《第七史诗》的付费点全部折算到一个基准单位 Φ（一次 UP 硬保底）上做横向比价，建立边际价值衰减与抽数时点敏感性两个模型，回答长尾付费为什么断掉、转化为什么集中出现在最后 40 抽。',
    'e7-payment-analysis': u'按结算货币把《第七史诗》商城拆成四轨商店，覆盖限时礼包、常驻直购、凭证兑换与订阅四类付费模块，含月免费投放总量、各付费深度资源比较与 121 抽保底下的限定角色获取周期演算。',
    'hsr-payment-analysis': u'以版本资源投放与五类用户的资源比较为主线拆解星铁付费体系：兑换比例与保底结构、月卡与通行证价值、各层用户获取限定角色的时长，延伸到限定角色定价与二手账号行情，给出可落地的商业化判断。',
    'wz-gift-analysis': u'梳理《王者荣耀》商城的礼包层级与定价逻辑，拆解首充、周卡、直购、限时四类礼包的投放节奏与锚定方式，并给出可迁移到其它 MOBA 手游的礼包设计清单。',
    'ares-combat-analysis': u'拆解《阿瑞斯病毒》的双摇杆战斗：从输入delay、受击帧、走位与弹药经济四个维度解释它的紧张感来源，并把移动端射击中的"笨重感"还原成一组可配置的参数。',
    'oasis-xiuxian-analysis': u'拆解《和平精英·绿洲启元》中「吃丹修仙」UGC 玩法的核心循环，还原境界突破—功法—资源投放三者的耦合关系，分析它在 UGC 平台上受欢迎的节奏设计原因。',
}


def estr(s):
    return html.escape(s, quote=False)


def heading_level(p, text, a_max, body):
    """判定段落是否为标题及其层级；0 表示正文。"""
    style = (p.style.name or '')
    m = re.search(r'(heading|标题)\s*([1-9])', style, re.I)
    if m:
        return min(int(m.group(2)) + 1, 4)
    sizes = [r.font.size.pt for r in p.runs if r.font.size]
    sz = max(sizes) if sizes else 0
    bold = any(r.bold for r in p.runs)
    if len(text) > 45:
        return 0
    if not (bold or (sz and sz > body + 1)):
        return 0
    if re.match(r'^(?:[一二三四五六七八九十]{1,3}|附)[、．.]', text):
        return 2
    if re.match(r'^\d{1,2}[、．]\s*\S', text):
        return 2
    if re.match(r'^\d{1,2}\.\d{1,2}\.\d{1,2}\b', text):
        return 4
    if re.match(r'^\d{1,2}\.\d{1,2}\b', text):
        return 3
    if re.match(r'^[（(][一二三四五六七八九十\d]{1,2}[）)]', text):
        return 3
    return 3 if sz > body else 0


def para_html(p, a_max, body):
    """段落 → HTML。返回 (html片段列表, 内嵌图片 rId 列表)"""
    rids = [b.get(qn('r:embed')) for b in p._p.findall('.//' + qn('a:blip'))]
    plain = p.text.strip()
    if not plain:
        return [], rids
    lvl = heading_level(p, plain, a_max, body)
    if lvl:
        return ['<h%d>%s</h%d>' % (lvl, estr(plain), lvl)], rids
    parts = []
    for run in p.runs:
        t = run.text
        if not t:
            continue
        if not t.strip():
            parts.append(' ')
            continue
        piece = estr(t)
        if run.bold:
            piece = '<strong>' + piece + '</strong>'
        elif run.italic:
            piece = '<em>' + piece + '</em>'
        parts.append(piece)
    text = ''.join(parts).strip()
    if not text:
        return [], rids
    return ['<p>' + text + '</p>'], rids


def collect_metrics(doc):
    """统计正文基础字号，以及加粗短段落里的最大字号，用于判定标题层级。"""
    sizes = {}
    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.size:
                sizes[r.font.size.pt] = sizes.get(r.font.size.pt, 0) + 1
    body = max(sizes.items(), key=lambda kv: kv[1])[0] if sizes else 10.5
    cand = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t or len(t) > 45:
            continue
        szs = [r.font.size.pt for r in p.runs if r.font.size]
        if szs and any(r.bold for r in p.runs):
            cand.append(max(szs))
    a_max = max(cand) if cand else body
    return a_max, body


def table_html(tb):
    rows = []
    for ri, row in enumerate(tb.rows):
        cells = []
        for c in row.cells:
            v = estr(c.text.strip().replace('\n', ' '))
            cells.append('<th>' + v + '</th>' if ri == 0 else '<td>' + v + '</td>')
        rows.append('<tr>' + ''.join(cells) + '</tr>')
    return '<div class="table-scroll"><table>' + ''.join(rows) + '</table></div>'


def docx_to_html(path, slug, skip_head=0):
    doc = Document(path)
    imgdir = os.path.join(IMG, slug)
    os.makedirs(imgdir, exist_ok=True)
    for f in os.listdir(imgdir):
        try:
            os.remove(os.path.join(imgdir, f))
        except OSError:
            pass
    rel = 'assets/img/' + slug + '/'
    a_max, body = collect_metrics(doc)
    parts = []
    img_seq = 0
    for child in doc.element.body.iterchildren():
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            p = Paragraph(child, doc)
            hs, rids = para_html(p, a_max, body)
            for rid in rids:
                try:
                    part = doc.part.related_parts[rid]
                except KeyError:
                    continue
                if 'image' not in part.content_type:
                    continue
                ext = part.content_type.split('/')[-1].replace('jpeg', 'jpg')
                img_seq += 1
                fname = 'fig%02d.%s' % (img_seq, ext)
                with open(os.path.join(imgdir, fname), 'wb') as fh:
                    fh.write(part.blob)
                parts.append('<figure class="figure"><img src="%s%s" alt="图 %d"><figcaption>图 %d</figcaption></figure>'
                             % (rel, fname, img_seq, img_seq))
            parts.extend(hs)
        elif tag == 'tbl':
            tb = Table(child, doc)
            parts.append(table_html(tb))
    body = '\n'.join(parts)
    body = body.replace(u'\ufffd', u'')
    for old_name, new_name in optimize_images(imgdir).items():
        body = body.replace(old_name, new_name)
    if skip_head:
        parts = body.split('\n')
        body = '\n'.join(parts[skip_head:])
    return strip_front(body)


CHAP = re.compile(r'<h2>(?:[一二三四五六七八九十附]+[、．.]|\d+[、．.])')


def strip_front(body):
    """丢掉封面区（文档大标题 / 作者 / 日期），从第一个章节号开始"""
    m = CHAP.search(body)
    if m and m.start() < 4000:
        return body[m.start():]
    return body


def md_to_html(path):
    src = open(path, encoding='utf-8').read()
    src = src.replace(u'\ufffd', u'')
    lines = src.split('\n')
    # 去掉文档自带的一级标题，避免与页面头部重复
    for i, ln in enumerate(lines[:8]):
        if ln.lstrip().startswith('# '):
            del lines[i]
            break
    src = '\n'.join(lines)
    body = markdown.markdown(src, extensions=['tables', 'fenced_code', 'sane_lists'])
    return body


def optimize_images(imgdir):
    """压缩图片并按需转 JPEG，返回 {原文件名: 新文件名} 便于回填 HTML"""
    renames = {}
    try:
        from PIL import Image
    except ImportError:
        print('   (未安装 pillow，跳过图片压缩)')
        return renames
    for name in sorted(os.listdir(imgdir)):
        fp = os.path.join(imgdir, name)
        try:
            im = Image.open(fp)
        except Exception:
            continue
        w, h = im.size
        if w > 1400:
            im = im.resize((1400, int(h * 1400.0 / w)), Image.LANCZOS)
        size_before = os.path.getsize(fp)
        if size_before > 130 * 1024:
            if im.mode in ('RGBA', 'LA', 'P'):
                bg = Image.new('RGB', im.size, (255, 255, 255))
                rgba = im.convert('RGBA')
                bg.paste(rgba, mask=rgba.split()[-1])
                im = bg
            else:
                im = im.convert('RGB')
            target = os.path.splitext(name)[0] + '.jpg'
            im.save(os.path.join(imgdir, target), format='JPEG', quality=78, optimize=True)
            if os.path.getsize(os.path.join(imgdir, target)) < size_before:
                os.remove(fp)
                renames[name] = target
                continue
            os.remove(os.path.join(imgdir, target))
        tmp = fp + '.tmp'
        im.save(tmp, format='PNG', optimize=True)
        if os.path.getsize(tmp) < size_before:
            os.replace(tmp, fp)
        else:
            os.remove(tmp)
    return renames


def build(work):
    if work['kind'] == 'md':
        if not os.path.exists(work['source']):
            print('!! 缺失 md:', work['source'])
            return None
        body = md_to_html(work['source'])
    else:
        if not os.path.exists(work['source']):
            print('!! 缺失 docx:', work['source'])
            return None
        body = docx_to_html(work['source'], work['slug'], work.get('skip_head', 0))
    if u'\ufffd' in body:
        body = body.replace(u'\ufffd', u'')
    htmlout = TPL.format(
        title=estr(work['title']),
        cat_label=work['cat_label'],
        tag=work['tag'],
        desc=estr(DESCS.get(work['slug'], u'')),
        role=estr(work['role']),
        year=work['year'],
        body=body,
    )
    outpath = os.path.join(OUT, work['slug'] + '.html')
    with open(outpath, 'w', encoding='utf-8') as fh:
        fh.write(htmlout)
    return outpath


def main():
    os.makedirs(OUT, exist_ok=True)
    total = 0
    for w in WORKS:
        p = build(w)
        if p:
            size = os.path.getsize(p)
            bad = u'\ufffd'
            raw = open(p, encoding='utf-8').read()
            print('%-26s %7.1f KB   乱码=%d   h2=%d   table=%d'
                  % (w['slug'], size / 1024.0, raw.count(bad),
                     len(re.findall(r'<h2', raw)), len(re.findall(r'<table', raw))))
            total += 1
    print('生成完毕：%d / %d' % (total, len(WORKS)))


if __name__ == '__main__':
    main()
